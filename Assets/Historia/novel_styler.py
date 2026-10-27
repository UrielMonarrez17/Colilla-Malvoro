import os
import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def create_novel_document():
    doc = docx.Document()
    
    # Configure 1-inch margins (2.54 cm)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        section.header_distance = Inches(0.5)
        section.footer_distance = Inches(0.5)

    return doc

def add_book_title(doc, title, subtitle=None, author="Colilla-Malvoro: Crónicas de la Niebla"):
    # Blank space before title
    p_pre = doc.add_paragraph()
    p_pre.paragraph_format.space_before = Pt(36)
    p_pre.paragraph_format.space_after = Pt(12)
    
    # Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_after = Pt(8)
    run_title = p_title.add_run(title)
    run_title.font.name = "Georgia"
    run_title.font.size = Pt(26)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(26, 36, 54) # Deep slate blue

    # Subtitle
    if subtitle:
        p_sub = doc.add_paragraph()
        p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_sub.paragraph_format.space_after = Pt(18)
        run_sub = p_sub.add_run(subtitle)
        run_sub.font.name = "Georgia"
        run_sub.font.size = Pt(13)
        run_sub.font.italic = True
        run_sub.font.color.rgb = RGBColor(90, 105, 125)

    # Decorative separator
    add_divider(doc, style="ornate")

    # Author / Metadata tag
    p_auth = doc.add_paragraph()
    p_auth.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_auth.paragraph_format.space_before = Pt(8)
    p_auth.paragraph_format.space_after = Pt(32)
    run_auth = p_auth.add_run(author)
    run_auth.font.name = "Calibri"
    run_auth.font.size = Pt(10)
    run_auth.font.bold = True
    run_auth.font.color.rgb = RGBColor(120, 120, 120)

def add_epigraph(doc, quote_text, source=None):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.8)
    p.paragraph_format.right_indent = Inches(0.8)
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(24)
    
    run_quote = p.add_run(f'«{quote_text}»\n')
    run_quote.font.name = "Georgia"
    run_quote.font.size = Pt(10.5)
    run_quote.font.italic = True
    run_quote.font.color.rgb = RGBColor(70, 70, 70)
    
    if source:
        run_src = p.add_run(f'— {source}')
        run_src.font.name = "Georgia"
        run_src.font.size = Pt(9.5)
        run_src.font.bold = True
        run_src.font.color.rgb = RGBColor(110, 110, 110)

def add_chapter_heading(doc, chapter_title, subtitle=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(28)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    
    run = p.add_run(chapter_title)
    run.font.name = "Georgia"
    run.font.size = Pt(16)
    run.font.bold = True
    run.font.color.rgb = RGBColor(20, 30, 45)

    if subtitle:
        p_sub = doc.add_paragraph()
        p_sub.paragraph_format.space_before = Pt(0)
        p_sub.paragraph_format.space_after = Pt(14)
        p_sub.paragraph_format.keep_with_next = True
        run_s = p_sub.add_run(subtitle)
        run_s.font.name = "Georgia"
        run_s.font.size = Pt(11.5)
        run_s.font.italic = True
        run_s.font.color.rgb = RGBColor(100, 110, 125)

def add_prose_paragraph(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0)
    p.paragraph_format.first_line_indent = Inches(0.25)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.15
    
    run = p.add_run(text)
    run.font.name = "Georgia"
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(35, 35, 35)

def add_dialogue(doc, speaker, speech, continuation=None):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.15
    
    if speaker:
        run_lead = p.add_run(f'—{speech}')
    else:
        run_lead = p.add_run(f'—{speech}')
        
    run_lead.font.name = "Georgia"
    run_lead.font.size = Pt(11)
    run_lead.font.color.rgb = RGBColor(30, 30, 30)

    if continuation:
        run_cont = p.add_run(f' {continuation}')
        run_cont.font.name = "Georgia"
        run_cont.font.size = Pt(11)
        run_cont.font.italic = True
        run_cont.font.color.rgb = RGBColor(70, 70, 70)

def add_divider(doc, style="simple"):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(12)
    
    divider_text = "✦  ✦  ✦" if style == "ornate" else "— · — · —"
    run = p.add_run(divider_text)
    run.font.name = "Georgia"
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(140, 150, 165)

def save_novel(doc, filename):
    # Save both in Assets/Historia and ../Historia
    paths = [
        os.path.join("Historia", filename),
        os.path.join("..", "Historia", filename)
    ]
    for p in paths:
        os.makedirs(os.path.dirname(os.path.abspath(p)), exist_ok=True)
        doc.save(p)
        print(f"Novela guardada exitosamente en: {os.path.abspath(p)}")
